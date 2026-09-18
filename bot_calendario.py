import re
from io import StringIO
from bs4 import BeautifulSoup
from curl_cffi import requests
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

print("=== INICIANDO EXTRACCIÓN CON FORMATO DE PLANTILLA Y FILTRO ===")

URL = "https://es.investing.com/economic-calendar/"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "es-ES,es;q=0.9",
}

# -------------------------------------------------------------
# 1. LISTA DE PAÍSES / DIVISAS PERMITIDOS (FILTRO)
# Agrega o quita los códigos que quieras incluir en tu Excel
# -------------------------------------------------------------
DIVISAS_PERMITIDAS = [
    "USD 🇺🇸",
    "EUR 🇪🇺",
    "EUR 🇩🇪",
    "EUR 🇫🇷",
    "GBP 🇬🇧",
    "CAD 🇨🇦",
    "JPY 🇯🇵",
    "CNY 🇨🇳",
    "NZD 🇳🇿",
    "AUD 🇦🇺",
    "CHF 🇨🇭",
]

MAPA_DIVISAS = {
    "US": "USD 🇺🇸",
    "USD": "USD 🇺🇸",
    "EUR": "EUR 🇪🇺",
    "DE": "EUR 🇩🇪",
    "FR": "EUR 🇫🇷",
    "IT": "EUR 🇮🇹",
    "ES": "EUR 🇪🇸",
    "GBP": "GBP 🇬🇧",
    "UK": "GBP 🇬🇧",
    "CAD": "CAD 🇨🇦",
    "AUD": "AUD 🇦🇺",
    "NZD": "NZD 🇳🇿",
    "JPY": "JPY 🇯🇵",
    "JP": "JPY 🇯🇵",
    "CNY": "CNY 🇨🇳",
    "CN": "CNY 🇨🇳",
    "CHF": "CHF 🇨🇭",
    "MXN": "MXN 🇲🇽",
    "BR": "BRL 🇧🇷",
}

# -------------------------------------------------------------
# ESTILOS DE EXCEL (Colores y Bordes para la foto izquierda)
# -------------------------------------------------------------
FILL_BANNER = PatternFill(
    start_color="1B4F72", end_color="1B4F72", fill_type="solid"
)  # Azul oscuro
FONT_BANNER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

FONT_HEADER = Font(name="Calibri", size=10, bold=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")

BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)

try:
    response = requests.get(
        URL, headers=HEADERS, impersonate="chrome124", timeout=15
    )

    if response.status_code == 200:
        soup = BeautifulSoup(response.text, "html.parser")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Histórico Eventos"

        # Nota superior informativa
        ws.merge_cells("B1:G1")
        cell_nota = ws["A1"]
        cell_nota.value = "Nota:"
        cell_nota.font = Font(bold=True)

        ws["B1"].value = (
            "Este archivo constituye la base histórica de indicadores macroeconómicos. "
            "Datos extraídos de investing.com, e intenta recopilar los eventos o indicadores publicados."
        )
        ws["B1"].font = Font(color="595959")

        filas_tabla = soup.find_all("tr")
        total_eventos = 0

        for tr in filas_tabla:
            classes = tr.get("class", [])

            # 2. ENCABEZADO DE FECHA (Banner Azul)
            if "theader" in classes or tr.find(
                "td", class_=lambda c: c and "theader" in c
            ):
                texto_fecha = tr.text.strip()
                if texto_fecha:
                    ws.append([])  # Fila vacía de separación

                    # Fila Banner Azul
                    current_row = ws.max_row + 1
                    ws.append([texto_fecha, "", "", "", "", "", ""])
                    ws.merge_cells(
                        start_row=current_row,
                        start_column=1,
                        end_row=current_row,
                        end_column=7,
                    )

                    cell_date = ws.cell(row=current_row, column=1)
                    cell_date.fill = FILL_BANNER
                    cell_date.font = FONT_BANNER
                    cell_date.alignment = ALIGN_CENTER

                    # Fila de Títulos de Columna
                    ws.append([
                        "Hora",
                        "Divisa",
                        "Impacto",
                        "Evento",
                        "Actual",
                        "Pronóstico",
                        "Anterior",
                    ])
                    header_row = ws.max_row
                    for col_idx in range(1, 8):
                        c = ws.cell(row=header_row, column=col_idx)
                        c.font = FONT_HEADER
                        c.alignment = ALIGN_CENTER
                        c.border = BORDER_THIN
                continue

            # 3. EXTRAER FILAS DE DATOS
            tds = tr.find_all("td")
            if len(tds) >= 6:
                # Hora
                hora = tds[0].text.strip()
                if not hora or hora.lower() == "hora":
                    continue

                # Divisa / País
                divisa_raw = tds[1].text.strip()
                flag_span = tds[1].find("span")
                if flag_span and flag_span.get("title"):
                    divisa_raw = flag_span.get("title").upper()
                divisa = MAPA_DIVISAS.get(divisa_raw, divisa_raw)

                # FILTRO DE PAÍSES: Si la divisa no está en la lista permitida, se ignora
                if DIVISAS_PERMITIDAS and divisa not in DIVISAS_PERMITIDAS:
                    continue

                # Impacto (Estrellas)
                bulls = tds[2].find_all(
                    "i", class_=re.compile(r"grayFullBullishIcon|grayBull")
                )
                num_stars = len(bulls) or tds[2].text.count("⭐") or 1
                impacto = "⭐" * num_stars

                # Evento
                evento_raw = tds[3].text.strip()
                evento = re.sub(
                    r"Act:.*", "", evento_raw, flags=re.IGNORECASE
                ).strip()

                # Valores
                actual = tds[4].text.strip() if len(tds) > 4 else "—"
                pronostico = tds[5].text.strip() if len(tds) > 5 else "—"
                anterior = tds[6].text.strip() if len(tds) > 6 else "—"

                # Escribir en Excel con bordes y alineaciones
                ws.append([
                    hora,
                    divisa,
                    impacto,
                    evento,
                    actual or "—",
                    pronostico or "—",
                    anterior or "—",
                ])
                row_idx = ws.max_row

                # Estilos para celdas de datos
                for col_idx in range(1, 8):
                    c = ws.cell(row=row_idx, column=col_idx)
                    c.border = BORDER_THIN
                    c.alignment = ALIGN_CENTER if col_idx != 4 else ALIGN_LEFT

                total_eventos += 1

        # Ajustar ancho automático de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        archivo_salida = "Historico_Eventos_Automatizado.xlsx"
        wb.save(archivo_salida)
        print(
            f"\n¡ÉXITO! Se generó el archivo con diseño idéntico y {total_eventos} eventos filtrados."
        )

    else:
        print(f"Error HTTP: {response.status_code}")

except Exception as e:
    print(f"Error en la ejecución: {e}")